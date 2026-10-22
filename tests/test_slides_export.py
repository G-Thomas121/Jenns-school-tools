"""
Turning a slides.md outline into a .pptx that Canva imports as an editable design.
"""
import io

import pytest
from pptx import Presentation
from pptx.util import Emu, Inches

from app.services import slides_export

DECK = """# Slides — *The Ramayana*, Canto 3

## Slide 1 — Bell Ringer
**Would you want to know the ending of a movie?**
Why might a storyteller tell you the whole story up front?

- Quick-write: 3 minutes
- Share with a *partner*
  - then share out

> Notes: Give 3 minutes of silent writing.
> Then take 3–4 responses.

## Slide 2 — Meet the Cast
| Heroes | Villains |
|---|---|
| **Rama**: the hero | **Ravan**: demon king |
| Sita | Kaikeyi |

> Notes: Don't spend long here.

## Exit Ticket
1. Define *argument*.
2. Name the hero.
"""


def pptx(markdown: str) -> Presentation:
    return Presentation(io.BytesIO(slides_export.to_pptx(markdown)))


def texts(slide) -> list[str]:
    return [p.text for s in slide.shapes if s.has_text_frame for p in s.text_frame.paragraphs if p.text]


# ── Parsing ────────────────────────────────────────────────────────────────────

def test_each_h2_is_a_slide_with_its_number_dropped_from_the_title():
    slides = slides_export.parse(DECK)
    assert [s.title for s in slides] == ["Bell Ringer", "Meet the Cast", "Exit Ticket"]


def test_notes_are_pulled_out_of_the_slide_body():
    first = slides_export.parse(DECK)[0]
    assert first.notes == "Give 3 minutes of silent writing.\nThen take 3–4 responses."
    assert not any("Notes" in str(block) for block in first.blocks)


def test_slides_md_without_any_slides_is_an_error():
    with pytest.raises(slides_export.SlidesError, match="## Slide"):
        slides_export.to_pptx("# Just a title\n\nSome text")


# ── The .pptx ──────────────────────────────────────────────────────────────────

def test_deck_is_widescreen_with_one_slide_per_section():
    deck = pptx(DECK)
    assert len(deck.slides) == 3
    assert (deck.slide_width, deck.slide_height) == (Inches(13.333), Inches(7.5))


def test_slide_text_is_plain_text_without_markdown_symbols():
    slide = pptx(DECK).slides[0]
    assert texts(slide) == [
        "Bell Ringer",
        "Would you want to know the ending of a movie?",
        "Why might a storyteller tell you the whole story up front?",
        "Quick-write: 3 minutes",
        "Share with a partner",
        "then share out",
    ]


def test_bold_and_italic_survive_as_formatting():
    slide = pptx(DECK).slides[0]
    runs = {r.text: r.font for s in slide.shapes if s.has_text_frame
            for p in s.text_frame.paragraphs for r in p.runs}
    assert runs["Would you want to know the ending of a movie?"].bold is True
    assert runs["partner"].italic is True


def test_bullets_are_real_bullets_with_nesting():
    slide = pptx(DECK).slides[0]
    paras = {p.text: p for s in slide.shapes if s.has_text_frame for p in s.text_frame.paragraphs}
    assert paras["Quick-write: 3 minutes"]._p.find(".//{*}buChar") is not None
    assert paras["then share out"].level == 1
    assert paras["Why might a storyteller tell you the whole story up front?"]._p.find(".//{*}buChar") is None


def test_numbered_lists_keep_their_numbers():
    assert texts(pptx(DECK).slides[2]) == ["Exit Ticket", "1. Define argument.", "2. Name the hero."]


def test_tables_become_editable_tables():
    slide = pptx(DECK).slides[1]
    tables = [s.table for s in slide.shapes if s.has_table]
    assert len(tables) == 1
    rows = [[c.text for c in row.cells] for row in tables[0].rows]
    assert rows == [["Heroes", "Villains"], ["Rama: the hero", "Ravan: demon king"], ["Sita", "Kaikeyi"]]


def test_notes_become_speaker_notes():
    deck = pptx(DECK)
    assert deck.slides[0].notes_slide.notes_text_frame.text == \
        "Give 3 minutes of silent writing.\nThen take 3–4 responses."
    assert not deck.slides[2].has_notes_slide


def test_busy_slides_get_smaller_text_and_stay_on_the_slide():
    sparse = "## A\n- One idea\n"
    busy = "## B\n" + "".join(f"- Point {i}: " + "a fairly long line of explanation " * 3 + "\n" for i in range(9))
    size = lambda deck: max(r.font.size for s in deck.slides[0].shapes if s.has_text_frame and s.top > Inches(1.5)
                            for p in s.text_frame.paragraphs for r in p.runs)
    assert size(pptx(busy)) < size(pptx(sparse))
    deck = pptx(busy)
    for shape in deck.slides[0].shapes:
        assert shape.top + shape.height <= deck.slide_height


def test_links_images_and_rules_are_handled_quietly():
    deck = pptx("## A\nSee [the text](https://example.com) ![pic](x.png)\n\n---\n\n`code`\n")
    assert texts(deck.slides[0]) == ["A", "See the text", "code"]


def test_table_columns_are_sized_to_their_content():
    deck = pptx("## Jigsaw\n| Group | Part | Lines |\n|---|---|---|\n"
                "| 1 | The Forest, the Kidnapping, and New Allies | 69–100 |\n| 2 | The Search for Sita | 101–132 |\n")
    table = next(s for s in deck.slides[0].shapes if s.has_table)
    group, part, lines = (c.width for c in table.table.columns)
    assert part > 2 * group
    assert part > 2 * lines
    assert group + part + lines == table.width


def test_narrow_columns_still_fit_their_longest_word():
    deck = pptx("## Jigsaw\n| Group | Part | Lines |\n|---|---|---|\n"
                "| 1 | The Forest, the Kidnapping, and New Allies, then the long road home | 101–132 |\n")
    table = next(s for s in deck.slides[0].shapes if s.has_table)
    lines = table.table.columns[2].width
    size = table.table.cell(1, 2).text_frame.paragraphs[0].runs[0].font.size
    assert lines >= len("101–132") * size / 1.7 + 2 * slides_export.CELL_MARGIN
